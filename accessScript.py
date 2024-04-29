import openpyxl
import pyodbc

# Function to read data from Excel file
def read_excel(file_path, sheet_name):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb[sheet_name]
    data = []

    for row in sheet.iter_rows(values_only=True):
        data.append(row)

    return data

# Function to write data to Access database
def write_to_access(data, access_file_path, table_name):
    conn_str = (
        r'DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};'
        f'DBQ={access_file_path};'
    )
    conn = pyodbc.connect(conn_str)
    cursor = conn.cursor()

    for row in data:
        placeholders = ','.join('?' * len(row))
        query = f'INSERT INTO {table_name} VALUES ({placeholders})'
        print("hi")
        cursor.execute(query, row)

    conn.commit()
    conn.close()

# Example usage
if __name__ == "__main__":
    excel_file = 'C:\\Users\\SABE15132\\Documents\\AccessProject.xlsx'  # Update with your Excel file path
    access_file = 'C:\\Users\\SABE15132\\Downloads\\Revit DB.mdb'  # Update with your Access database file path
    sheet_name = 'Sheet1'  # Update with your sheet name
    table_name = 'Walls'  # Update with your table name

    data = read_excel(excel_file, sheet_name)
    print(data)
    write_to_access(data, access_file, table_name)